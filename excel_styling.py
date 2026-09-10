import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def style_my_excel_sheet():
    print("🚀 Starting Sudhir's Advanced Excel Styling Script...\n")
    
    excel_filename = "laptops.xlsx"
    
    try:
        # 1. Excel. Loading the file (which we created on Saturday)
        wb = openpyxl.load_workbook(excel_filename)
        ws = wb.active
        
        # 2.Creating styling designs  (Professional Navy Blue & White)
        header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Navy blue
        center_alignment = Alignment(horizontal="center", vertical="center")
        
        # Thin border design
        thin_side = Side(border_style="thin", color="D9D9D9")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        # 3.Formatting the first row (headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = cell_border
            
        # 4.Applying borders and alignment to the remaining data cells.
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.font = Font(name="Arial", size=11)
                cell.border = cell_border
                if cell.column == 1: #To center the serial number
                    cell.alignment = center_alignment
                    
        # 5. (Auto-fit Column Width)
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        # Saving the file again securely
        wb.save(excel_filename)
        print("🎉 Success! Your Excel sheet 'laptops.xlsx' is now beautifully formatted.")
        print("➡️ Open the file to see the magical Corporate Navy Blue layout!")
        
    except Exception as e:
        print(f"❌ Error occurred: {e}\n(Note: Make sure laptops.xlsx is closed before running)")

if __name__ == "__main__":
    style_my_excel_sheet()